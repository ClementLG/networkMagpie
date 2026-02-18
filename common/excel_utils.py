import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# --- Excel Configuration ---
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
BLUE_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
BOLD_FONT = Font(bold=True)

def apply_header_style(ws, row_num=1):
    """
    Applies a standard bold, blue styling to the header row of an Excel worksheet.

    Args:
        ws (openpyxl.worksheet.worksheet.Worksheet): The worksheet to apply styling to.
        row_num (int, optional): The row number to style as header. Defaults to 1.
    """
    for cell in ws[row_num]:
        cell.font = BOLD_FONT
        cell.fill = BLUE_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

def auto_fit_columns(ws):
    """
    Automatically adjusts the width of all columns in an Excel worksheet based on their content.

    Args:
        ws (openpyxl.worksheet.worksheet.Worksheet): The worksheet to adjust.
    """
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

def set_cell_status_color(cell, level):
    """
    Sets the background color of an Excel cell based on the security level/status.

    Args:
        cell (openpyxl.cell.cell.Cell): The cell to colorize.
        level (str): The security level ('good', 'warning', 'bad', 'error').
    """
    if level == "good":
        cell.fill = GREEN_FILL
    elif level == "warning":
        cell.fill = ORANGE_FILL
    elif level == "error" or level == "bad":
        cell.fill = RED_FILL
