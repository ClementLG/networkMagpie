import openpyxl
import logging
from .excel_utils import apply_header_style, auto_fit_columns, set_cell_status_color, GREEN_FILL, ORANGE_FILL, RED_FILL

logger = logging.getLogger(__name__)

def generate_excel_report(all_data, excel_filepath):
    """
    Generates a comprehensive Excel report from the collected audit data.

    Args:
        all_data (list): A list of dictionaries containing audit data for all devices.
        excel_filepath (str): The file path where the Excel report will be saved.
    """
    if not all_data:
        logger.warning("No data for Excel report.")
        return
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    # 1. General Info
    ws_info = wb.create_sheet("General Info")
    headers_info = ["Hostname", "IP Address", "Model", "OS Version", "Uptime", "Serial Number"]
    ws_info.append(headers_info)
    apply_header_style(ws_info)
    for dev_data in all_data:
        if dev_data.get('status') == 'error_connection':
            ws_info.append(
                [dev_data.get('attempted_host', 'N/A'), dev_data.get('attempted_host', 'N/A'), "CONNECTION ERROR",
                 dev_data.get('error_message', 'N/A'), "N/A", "N/A"])
            for i in range(3, 5): ws_info.cell(row=ws_info.max_row, column=i).fill = RED_FILL
        else:
            info = dev_data.get("general_info", {})
            ws_info.append([info.get("hostname", dev_data.get("host", "N/A")), dev_data.get("host", "N/A"),
                            info.get("model", "N/A"), info.get("ios_version", "N/A"), info.get("uptime", "N/A"),
                            info.get("serial_number", "N/A")])
    auto_fit_columns(ws_info)
    
    # 2. Interfaces
    ws_interfaces = wb.create_sheet("Interfaces")
    headers_interfaces = ["Hostname", "Interface", "Type", "Description", "IP Address", "Link Status",
                          "Protocol Status", "VLAN (Access)", "Duplex", "Speed"]
    ws_interfaces.append(headers_interfaces)
    apply_header_style(ws_interfaces)
    for dev_data in all_data:
        if dev_data.get('status') == 'error_connection': continue
        hostname = dev_data.get("general_info", {}).get("hostname", dev_data.get("host", "N/A"))
        for iface in dev_data.get("interfaces", []):
            ws_interfaces.append(
                [hostname, iface.get("name"), iface.get("type"), iface.get("description"), iface.get("ip_address"),
                 iface.get("status_link"), iface.get("status_protocol"), iface.get("vlan"), iface.get("duplex"),
                 iface.get("speed")])
            lr, pr = ws_interfaces.cell(row=ws_interfaces.max_row, column=6), ws_interfaces.cell(
                row=ws_interfaces.max_row, column=7)
            sl, sp = str(iface.get("status_link", "")).lower(), str(iface.get("status_protocol", "")).lower()
            if sl == "up" or "connected" in sl:
                lr.fill = GREEN_FILL
            elif ("admin" in sl and "down" in sl) or sl == "disabled" or "admin-down" in sl:
                lr.fill = ORANGE_FILL
            elif sl == "down" or "notconnect" in sl or "err-disabled" in sl or "error-disabled" in sl or "down (no transceiver)" in sl:
                lr.fill = RED_FILL
            if sp == "up":
                pr.fill = GREEN_FILL
            elif sp == "down":
                pr.fill = RED_FILL
    auto_fit_columns(ws_interfaces)
    
    # 3. VLANs
    ws_vlans = wb.create_sheet("VLANs")
    headers_vlans = ["Hostname", "VLAN ID", "VLAN Name", "Status", "Assigned Ports"]
    ws_vlans.append(headers_vlans)
    apply_header_style(ws_vlans)
    for dev_data in all_data:
        if dev_data.get('status') == 'error_connection': continue
        hostname = dev_data.get("general_info", {}).get("hostname", dev_data.get("host", "N/A"))
        for vlan in dev_data.get("vlans", []):
            ws_vlans.append([hostname, vlan.get("id"), vlan.get("name"), vlan.get("status"), vlan.get("ports")])
            sc = ws_vlans.cell(row=ws_vlans.max_row, column=4)
            if str(vlan.get("status", "")).lower() == "active" or str(vlan.get("status", "")).lower() == "up":
                sc.fill = GREEN_FILL
            else:
                sc.fill = ORANGE_FILL
    auto_fit_columns(ws_vlans)
    
    # 4. ARP Table
    ws_arp = wb.create_sheet("ARP Table")
    headers_arp = ["Hostname", "Protocol", "IP Address", "Age (min)", "MAC Address", "Type", "Interface"]
    ws_arp.append(headers_arp)
    apply_header_style(ws_arp)
    for dev_data in all_data:
        if dev_data.get('status') == 'error_connection': continue
        hostname = dev_data.get("general_info", {}).get("hostname", dev_data.get("host", "N/A"))
        for entry in dev_data.get("arp_table", []):
            ws_arp.append(
                [hostname, entry.get("protocol"), entry.get("address"), entry.get("age"), entry.get("mac_address"),
                 entry.get("type"), entry.get("interface")])
    auto_fit_columns(ws_arp)
    
    # 5. Security Audit
    ws_security = wb.create_sheet("Security Audit")
    headers_security = ["Hostname", "Check Name", "Status/Value", "Level", "Details/Recommendation"]
    ws_security.append(headers_security)
    apply_header_style(ws_security)
    for dev_data in all_data:
        if dev_data.get('status') == 'error_connection': continue
        hostname = dev_data.get("general_info", {}).get("hostname", dev_data.get("host", "N/A"))
        for check_name, check_data in dev_data.get("security_audit", {}).items():
            ws_security.append(
                [hostname, check_name.replace("_", " ").title(), str(check_data.get("status")), check_data.get("level"),
                 check_data.get("details")])
            lc = ws_security.cell(row=ws_security.max_row, column=4)
            set_cell_status_color(lc, check_data.get("level"))
    auto_fit_columns(ws_security)
    
    try:
        wb.save(excel_filepath)
        logger.info(f"[+] Excel report generated: {excel_filepath}")
    except Exception as e:
        logger.error(f"[-] Error saving Excel: {e}")
